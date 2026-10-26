# -*- coding: utf-8 -*-
"""Paginate Project_Report_VERITAS.docx with Word and sync TOC / LoT / LoF page numbers."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Project_Report_VERITAS.docx"
MAP = ROOT / "report_figures" / "page_map.json"


def clear_and_fill_table(table, rows_data: list[list[str]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    hdr = table.rows[0].cells
    for i, val in enumerate(rows_data[0]):
        if i < len(hdr):
            hdr[i].text = val
    for row in rows_data[1:]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if i < len(cells):
                cells[i].text = val


def word_page_map(docx_path: Path) -> dict:
    import win32com.client  # type: ignore

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    items = []
    try:
        doc = word.Documents.Open(str(docx_path), False, True)
        doc.Repaginate()
        total = int(doc.ComputeStatistics(2))
        for p in doc.Paragraphs:
            t = (p.Range.Text or "").replace("\r", "").replace("\x07", "").strip()
            if not t:
                continue
            style = p.Style.NameLocal
            keep = (
                style.startswith("Heading")
                or t.startswith(("Fig ", "Figure ", "Table ", "ABSTRACT", "LIST OF", "TABLE OF"))
                or (re.match(r"^\d{1,2}\. ", t) is not None and len(t) < 160)
            )
            if not keep:
                continue
            page = int(p.Range.Information(1))
            items.append({"text": t[:160], "style": style, "page": page})
        doc.Close(False)
    finally:
        word.Quit()
    payload = {"total_pages": total, "items": items}
    MAP.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return payload


def find_page(items: list[dict], prefixes: tuple[str, ...]) -> int | None:
    for it in items:
        t = it["text"]
        for pref in prefixes:
            if t == pref or t.startswith(pref):
                return int(it["page"])
    return None


def roman(n: int) -> str:
    vals = [
        (10, "x"),
        (9, "ix"),
        (5, "v"),
        (4, "iv"),
        (1, "i"),
    ]
    out = []
    for v, s in vals:
        while n >= v:
            out.append(s)
            n -= v
    return "".join(out)


def sync_lists(docx_path: Path, page_map: dict) -> None:
    items = page_map["items"]

    def p(*prefs: str, default: str = "?") -> str:
        n = find_page(items, prefs)
        return str(n) if n is not None else default

    # Front matter absolute pages → Roman (body Arabic restarts at 1)
    lot_front = find_page(items, ("LIST OF TABLES",))
    lof_front = find_page(items, ("LIST OF FIGURES",))
    abs_front = find_page(items, ("ABSTRACT",))
    # Prefer later occurrences (actual headings after TOC labels)
    lot_pages = [it["page"] for it in items if it["text"].startswith("LIST OF TABLES")]
    lof_pages = [it["page"] for it in items if it["text"].startswith("LIST OF FIGURES")]
    abs_pages = [it["page"] for it in items if it["text"].startswith("ABSTRACT")]
    lot_r = roman(lot_pages[-1]) if lot_pages else "v"
    lof_r = roman(lof_pages[-1]) if lof_pages else "vi"
    abs_r = roman(abs_pages[-1]) if abs_pages else "vii"

    toc = [
        ["CHAPTER NO.", "TITLE", "PAGE NO."],
        ["", "LIST OF TABLES", lot_r],
        ["", "LIST OF FIGURES", lof_r],
        ["", "ABSTRACT", abs_r],
        ["1", "INTRODUCTION", p("1. INTRODUCTION")],
        ["", "1.1 Background", p("1.1 Background")],
        ["", "1.2 Problem Statement", p("1.2 Problem Statement")],
        ["", "1.3 Motivation", p("1.3 Motivation")],
        ["", "1.4 Objectives", p("1.4 Objectives")],
        ["", "1.5 Scope", p("1.5 Scope")],
        ["", "1.6 Justification", p("1.6 Justification")],
        ["2", "LITERATURE SURVEY", p("2. LITERATURE SURVEY")],
        ["3", "EXISTING SYSTEM", p("3. EXISTING SYSTEM")],
        ["4", "PROPOSED SYSTEM", p("4. PROPOSED SYSTEM")],
        ["5", "SYSTEM ARCHITECTURE", p("5. SYSTEM ARCHITECTURE")],
        ["6", "MODULE DESCRIPTION", p("6. MODULE DESCRIPTION")],
        ["7", "THREAT MODEL", p("7. THREAT MODEL")],
        ["8", "VULNERABILITY ANALYSIS", p("8. VULNERABILITY ANALYSIS")],
        ["9", "RISK ASSESSMENT", p("9. RISK ASSESSMENT")],
        ["10", "IMPLEMENTATION", p("10. IMPLEMENTATION")],
        ["11", "SECURITY IMPLEMENTATION", p("11. SECURITY IMPLEMENTATION")],
        ["12", "SECURITY TESTING", p("12. SECURITY TESTING")],
        ["13", "RESULTS AND ANALYSIS", p("13. RESULTS AND ANALYSIS")],
        ["14", "CONCLUSION", p("14. CONCLUSION")],
        ["15", "FUTURE ENHANCEMENT", p("15. FUTURE ENHANCEMENT")],
        ["16", "REFERENCES", p("16. REFERENCES")],
    ]

    lot = [
        ["TABLE NO.", "TITLE", "PAGE NO."],
        ["2.1", "Literature Survey", p("Table 2.1:")],
        ["3.1", "Existing System Comparison", p("Table 3.1:")],
        ["7.1", "Threats Affecting Major Assets", p("Table 7.1:")],
        ["8.1", "Vulnerability Analysis", p("Table 8.1:")],
        ["9.1", "Risk Register", p("Table 9.1:")],
        ["12.1", "Security Test Cases", p("Table 12.1:")],
        ["13.1", "Results Table / Metrics", p("Table 13.1:")],
    ]

    lof = [
        ["FIG. NO.", "TITLE", "PAGE NO."],
        ["4.1", "Proposed Workflow", p("Fig 4.1:")],
        ["5.1", "System Architecture", p("Fig 5.1:")],
        ["5.2", "Data Flow Diagram (Level 1)", p("Fig 5.2:")],
        ["10.1", "User Login Interface (live VERITAS)", p("Fig 10.1:")],
        ["10.2", "User Projects Workspace (live production VERITAS)", p("Fig 10.2:")],
        ["10.3", "Administrator Dashboard (live VERITAS)", p("Fig 10.3:")],
        ["13.1", "Live Production Scan Findings — NurseQuest", p("Fig 13.1:")],
        ["13.2", "Live Production Home — Findings by Run and Severity", p("Fig 13.2:")],
        ["13.3", "Live Production Scan Findings — AI Honeypot", p("Fig 13.3:")],
    ]

    doc = Document(str(docx_path))
    clear_and_fill_table(doc.tables[1], toc)
    clear_and_fill_table(doc.tables[2], lot)
    clear_and_fill_table(doc.tables[3], lof)
    doc.save(str(docx_path))

    print("Synced TOC / List of Tables / List of Figures")
    print("TOC sample:")
    for row in toc[1:]:
        print(" ", " | ".join(row))
    print("LoT:")
    for row in lot[1:]:
        print(" ", " | ".join(row))
    print("LoF:")
    for row in lof[1:]:
        print(" ", " | ".join(row))


def main() -> None:
    if not DOCX.exists():
        raise SystemExit(f"Missing {DOCX}")
    print(f"Paginating {DOCX} with Word…")
    try:
        page_map = word_page_map(DOCX)
    except Exception as e:
        print(f"win32com failed ({e}); trying PowerShell COM fallback is required", file=sys.stderr)
        raise
    print(f"Total pages: {page_map['total_pages']}, mapped items: {len(page_map['items'])}")
    sync_lists(DOCX, page_map)


if __name__ == "__main__":
    main()
