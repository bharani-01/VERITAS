from pptx import Presentation

path = r'D:\College\Sem-5\Computer Network\CN-Project\Report\NetGuard_Project_Presentation.pptx'
prs = Presentation(path)

print(f"Total slides: {len(prs.slides)}")
for i, s in enumerate(prs.slides):
    tag = ""
    title = ""
    subtitle = ""
    has_table = any(shape.has_table for shape in s.shapes)
    has_pic = any(shape.shape_type == 13 for shape in s.shapes)
    rects = sum(1 for shape in s.shapes if shape.shape_type == 1)
    
    # Try finding title/tag
    for shape in s.shapes:
        if shape.has_text_frame and shape.text_frame.text:
            lines = [l.strip() for l in shape.text_frame.text.split('\n') if l.strip()]
            if lines and not title:
                if len(lines) >= 3:
                    tag = lines[0]
                    title = lines[1]
                    subtitle = lines[2]
                elif len(lines) >= 2:
                    title = lines[0]
                    subtitle = lines[1]
                else:
                    title = lines[0]
    print(f"Slide {i+1}: Tag='{tag}' | Title='{title}' | Subtitle='{subtitle}' | Table={has_table} | Pic={has_pic} | Rects={rects}")
