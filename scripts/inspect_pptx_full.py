import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = pptx.Presentation(r'c:\Users\Swaminathan\Downloads\NetGuard_Project_Presentation.pptx')

print(f"Slide width: {prs.slide_width / Inches(1):.2f} in, height: {prs.slide_height / Inches(1):.2f} in")
print(f"Total slides: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n--- SLIDE {idx} ---")
    print(f"Layout: {slide.slide_layout.name}")
    for shape in slide.shapes:
        name = shape.name
        stype = shape.shape_type
        left = shape.left / Inches(1)
        top = shape.top / Inches(1)
        width = shape.width / Inches(1)
        height = shape.height / Inches(1)
        print(f"  Shape: '{name}' | Type: {stype} | Bounds: ({left:.2f}, {top:.2f}, {width:.2f}, {height:.2f})")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                p_text = "".join(run.text for run in p.runs).strip()
                if p_text:
                    runs_info = []
                    for r in p.runs:
                        font_name = r.font.name
                        font_size = f"{r.font.size.pt:.1f}pt" if r.font.size else "default"
                        bold = "B" if r.font.bold else ""
                        color = r.font.color.rgb if r.font.color and r.font.color.type == 1 else "scheme/none"
                        runs_info.append(f"'{r.text}' ({font_name}, {font_size}, {bold}, {color})")
                    print(f"    P{p_idx} [{p.alignment}]: {' + '.join(runs_info)}")
        elif shape.has_table:
            table = shape.table
            print(f"    Table: {len(table.rows)} rows x {len(table.columns)} cols")
            for r_idx, row in enumerate(table.rows):
                row_vals = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                print(f"      R{r_idx}: {' | '.join(row_vals)}")
